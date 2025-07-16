"""
Excel file parser for AI Memory Bank
"""

from pathlib import Path
from typing import List, Dict, Any
import logging
import openpyxl
import pandas as pd
from config import CHUNK_SIZE, CHUNK_OVERLAP

logger = logging.getLogger(__name__)

def parse_excel(file_path: Path) -> List[Dict[str, Any]]:
    """
    Parse an Excel file and return structured content
    
    Args:
        file_path: Path to the Excel file
    
    Returns:
        List of dictionaries containing parsed content and metadata
    """
    try:
        # Load the workbook
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        
        all_content = []
        sheet_info = []
        
        for sheet_name in workbook.sheetnames:
            try:
                sheet = workbook[sheet_name]
                
                # Convert sheet to DataFrame for easier processing
                data = []
                for row in sheet.iter_rows(values_only=True):
                    data.append(row)
                
                df = pd.DataFrame(data)
                
                # Remove completely empty rows and columns
                df = df.dropna(how='all').dropna(axis=1, how='all')
                
                if not df.empty:
                    # Convert DataFrame to text representation
                    sheet_text = _dataframe_to_text(df, sheet_name)
                    sheet_info.append({
                        "name": sheet_name,
                        "rows": len(df),
                        "columns": len(df.columns),
                        "text": sheet_text
                    })
                    all_content.append(sheet_text)
                
            except Exception as e:
                logger.warning(f"Failed to parse sheet '{sheet_name}': {e}")
                continue
        
        # Combine all content
        combined_text = "\n\n".join(all_content)
        
        # Clean the text
        combined_text = _clean_excel_text(combined_text)
        
        # Split into chunks
        chunks = _chunk_excel_text(combined_text, sheet_info)
        
        # Create structured output
        results = []
        for i, chunk in enumerate(chunks):
            if chunk["content"].strip():  # Skip empty chunks
                results.append({
                    "content": chunk["content"],
                    "metadata": {
                        "file_path": str(file_path),
                        "file_type": "excel",
                        "chunk_index": i,
                        "total_chunks": len(chunks),
                        "chunk_size": len(chunk["content"]),
                        "file_size": len(combined_text),
                        "total_sheets": len(workbook.sheetnames),
                        "sheets_in_chunk": chunk.get("sheets", []),
                        "excel_info": {
                            "workbook_properties": {
                                "title": workbook.properties.title or "",
                                "creator": workbook.properties.creator or "",
                                "subject": workbook.properties.subject or "",
                                "keywords": workbook.properties.keywords or ""
                            }
                        }
                    }
                })
        
        logger.info(f"Parsed {file_path}: {len(results)} chunks created from {len(workbook.sheetnames)} sheets")
        return results
        
    except Exception as e:
        logger.error(f"Failed to parse {file_path}: {e}")
        raise

def _dataframe_to_text(df: pd.DataFrame, sheet_name: str) -> str:
    """Convert DataFrame to text representation"""
    text_parts = [f"Sheet: {sheet_name}"]
    text_parts.append("=" * (len(sheet_name) + 7))
    
    # Add column headers if they exist
    if not df.empty:
        # Check if first row looks like headers
        first_row = df.iloc[0]
        if any(isinstance(cell, str) and cell.strip() for cell in first_row):
            headers = [str(cell) if cell is not None else f"Column_{i}" for i, cell in enumerate(first_row)]
            df.columns = headers
            df = df.iloc[1:]  # Remove header row from data
        
        # Convert to string representation
        text_parts.append(df.to_string(index=False, na_rep=''))
    
    return "\n".join(text_parts)

def _clean_excel_text(text: str) -> str:
    """Clean and normalize Excel text content"""
    import re
    
    # Remove excessive whitespace
    text = re.sub(r'\s+', ' ', text)
    
    # Remove null characters
    text = text.replace('\x00', '')
    
    # Clean up sheet separators
    text = re.sub(r'Sheet: .*\n=+\n', '\n\n', text)
    
    # Normalize whitespace
    text = ' '.join(text.split())
    
    return text.strip()

def _chunk_excel_text(text: str, sheet_info: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Split Excel text into overlapping chunks, preserving sheet information
    
    Args:
        text: Full text content
        sheet_info: List of sheet information
    
    Returns:
        List of chunk dictionaries with content and sheet information
    """
    if len(text) <= CHUNK_SIZE:
        return [{"content": text, "sheets": [si["name"] for si in sheet_info]}]
    
    chunks = []
    current_chunk = ""
    start_pos = 0
    
    # Create a mapping of text positions to sheets
    sheet_positions = _map_text_to_sheets(text, sheet_info)
    
    while start_pos < len(text):
        end_pos = start_pos + CHUNK_SIZE
        
        # Try to break at sentence boundaries
        if end_pos < len(text):
            for i in range(end_pos, max(start_pos + CHUNK_SIZE - 100, start_pos), -1):
                if text[i] in '.!?':
                    end_pos = i + 1
                    break
        
        chunk_text = text[start_pos:end_pos].strip()
        
        if chunk_text:
            # Find which sheets this chunk contains
            chunk_sheets = set()
            for pos in range(start_pos, end_pos):
                if pos in sheet_positions:
                    chunk_sheets.update(sheet_positions[pos])
            
            chunks.append({
                "content": chunk_text,
                "sheets": sorted(list(chunk_sheets))
            })
        
        # Move start position with overlap
        start_pos = end_pos - CHUNK_OVERLAP
        if start_pos >= len(text):
            break
    
    return chunks

def _map_text_to_sheets(text: str, sheet_info: List[Dict[str, Any]]) -> Dict[int, List[str]]:
    """
    Create a mapping from text positions to sheet names
    
    Args:
        text: Full text content
        sheet_info: List of sheet information
    
    Returns:
        Dictionary mapping text positions to list of sheet names
    """
    position_to_sheets = {}
    current_pos = 0
    
    for sheet in sheet_info:
        sheet_text = sheet["text"]
        sheet_name = sheet["name"]
        
        # Find where this sheet text appears in the full text
        sheet_start = text.find(sheet_text, current_pos)
        if sheet_start != -1:
            sheet_end = sheet_start + len(sheet_text)
            
            # Mark all positions in this sheet
            for pos in range(sheet_start, sheet_end):
                if pos not in position_to_sheets:
                    position_to_sheets[pos] = []
                position_to_sheets[pos].append(sheet_name)
            
            current_pos = sheet_end
    
    return position_to_sheets 